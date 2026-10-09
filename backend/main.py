import io
import os
from collections import defaultdict
from datetime import date, datetime
from typing import List, Optional

from fastapi import Depends, FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel
from sqlalchemy import extract, func
from sqlalchemy.orm import Session

import models
from auth import UsuarioAtual, criar_token, usuarios_portal, verificar_senha
from database import Base, engine, get_db
from importer import ler_relatorio

app = FastAPI(title="Piaseg · Apólices canceladas")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])

STATUS_OPCOES = [
    "Cliente refez o seguro com a Piaseg (nova apólice/seguradora)",
    "Em negociação para recuperar o cliente",
    "Cliente vendeu o bem / encerrou a atividade",
    "Falta de pagamento (inadimplência)",
    "Cliente achou caro / reduziu custos",
    "Migrou para outra corretora",
    "Insatisfação com a seguradora ou atendimento",
    "Cancelado pela seguradora",
    "Sem retorno do cliente",
    "Outro motivo",
]

# Vínculo inicial login do portal -> franqueado (unidade de negócio), por semelhança de nome.
# Mesma lista do painel de não renovadas. Só é aplicado quando a tabela de vínculos está vazia;
# depois o gestor ajusta na tela "Acessos".
VINCULOS_INICIAIS = {
    "AGUILERA": ["thiago@piaseg.com.br"],
    "ALENCAR VEICULOS": ["leandrol.alencar@gmail.com", "rafael_alencar_@hotmail.com"],
    "ALEXANDRE SILVA AMORIM": ["alexandre.amorim@piaseg.com.br"],
    "ARRUDA & OURIQUES": ["energiasol"],
    "BRUNO LOURENCO GIROTTO": ["nathalha.mcg@gmail.com"],
    "CAMBARA SEGUROS": ["cambarasegurosenegocios@gmail.com"],
    "CELSO CARLOS CAVALLIERI JUNIOR": ["celsocarloscavallierijunior@gmail.com"],
    "DAYWID WILLIAM": ["daywid.wst@hotmail.com"],
    "FABIANA TEIXEIRA": ["fabianaseg2024@gmail.com"],
    "FERNANDO SILVA": ["fernandoscfilho2008@gmail.com"],
    "ILSON GOMES DA SILVA": ["ilsongomesh@hotmail.com"],
    "INFINITY SEG": ["infinitysegg@hotmail.com"],
    "IVAN / ROQUE": ["souvencer50@gmail.com"],
    "JIMMY CRISTIAN ALEGRE": ["jimmydourados@gmail.com"],
    "JOAO A P BOEIRA": ["joaoa.boeira@hotmail.com"],
    "LJKL NEGOCIOS": ["lucasjikal777@gmail.com"],
    "MARVSEGUROS": ["marvsegquiri@hotmail.com"],
    "MAURICIO PARDINHO": ["mauriciocg@piaseg.com.br"],
    "RENATA PIRES AMOROSO LIMA": ["renata.amoroso@piaseg.com.br"],
    "ROQUE HOSANO DOS SANTOS CRUZ": ["souvencer50@gmail.com"],
    "SONIMAR MACHADO": ["sonimarmachado@gmail.com"],
    "TERRA": ["seguros2@grupoterradourados.com.br", "seguro@imobiliariaterradourados.com.br"],
}
ADMINS_INICIAIS = ["admin"]


@app.on_event("startup")
def startup():
    Base.metadata.create_all(bind=engine)
    with Session(engine) as db:
        if db.query(models.Vinculo).count() == 0:
            portal = {}
            try:
                portal = {str(u.get("u", "")).lower(): u for u in usuarios_portal()}
            except Exception:
                pass
            for franqueado, usuarios in VINCULOS_INICIAIS.items():
                for u in usuarios:
                    # o login "energiasol…" está abreviado acima: resolve pelo prefixo
                    real = next((k for k in portal if k.startswith(u)), u)
                    db.add(models.Vinculo(usuario=real, franqueado=franqueado))
        if db.query(models.Admin).count() == 0:
            for u in ADMINS_INICIAIS:
                db.add(models.Admin(usuario=u))
        db.commit()


# ---------------------------------------------------------------- helpers

def eh_admin(db: Session, usuario: str) -> bool:
    return db.get(models.Admin, usuario) is not None


def franqueados_do_usuario(db: Session, usuario: str) -> List[str]:
    return sorted({v.franqueado for v in db.query(models.Vinculo).filter(models.Vinculo.usuario == usuario)})


def perfil(db: Session, user: dict) -> dict:
    admin = eh_admin(db, user["usuario"])
    return {
        "usuario": user["usuario"],
        "nome": user["nome"],
        "admin": admin,
        "franqueados": franqueados_do_usuario(db, user["usuario"]),
    }


def escopo(db: Session, user: dict, franqueado: Optional[str]):
    """Query de apólices limitada ao que o usuário pode ver."""
    q = db.query(models.Apolice)
    if eh_admin(db, user["usuario"]):
        if franqueado:
            q = q.filter(models.Apolice.franqueado == franqueado)
        return q
    meus = franqueados_do_usuario(db, user["usuario"])
    if franqueado:
        if franqueado not in meus:
            raise HTTPException(403, "Você não tem acesso a este franqueado.")
        meus = [franqueado]
    return q.filter(models.Apolice.franqueado.in_(meus or ["__nenhum__"]))


def apolice_json(a: models.Apolice) -> dict:
    return {
        "id": a.id,
        "franqueado": a.franqueado,
        "unidade": a.unidade,
        "cliente": a.cliente,
        "seguradora": a.seguradora,
        "cancelamento": a.cancelamento.isoformat(),
        "vigencia_inicio": a.vigencia_inicio.isoformat() if a.vigencia_inicio else None,
        "vigencia_fim": a.vigencia_fim.isoformat() if a.vigencia_fim else None,
        "numero": a.numero,
        "produto": a.produto,
        "premio": a.premio,
        "status": a.status,
        "respondido_em": a.respondido_em.isoformat() + "Z" if a.respondido_em else None,
        "no_ultimo_relatorio": a.no_ultimo_relatorio,
    }


def filtrar(q, ano: Optional[int], mes: Optional[int], situacao: Optional[str], busca: Optional[str]):
    if ano:
        q = q.filter(extract("year", models.Apolice.cancelamento) == ano)
    if mes:
        q = q.filter(extract("month", models.Apolice.cancelamento) == mes)
    if situacao == "pendente":
        q = q.filter(models.Apolice.status.is_(None))
    elif situacao == "respondida":
        q = q.filter(models.Apolice.status.isnot(None))
    if busca:
        termo = f"%{busca.strip()}%"
        q = q.filter(
            models.Apolice.cliente.ilike(termo)
            | models.Apolice.numero.ilike(termo)
            | models.Apolice.produto.ilike(termo)
            | models.Apolice.seguradora.ilike(termo)
        )
    return q


# ---------------------------------------------------------------- auth

class LoginIn(BaseModel):
    usuario: str
    senha: str


@app.get("/")
def raiz():
    return {"ok": True, "app": "apolices-canceladas"}


@app.post("/auth/login")
def login(dados: LoginIn, db: Session = Depends(get_db)):
    user = verificar_senha(dados.usuario, dados.senha)
    if not user:
        raise HTTPException(401, "Usuário ou senha incorretos. Use o mesmo acesso do Portal do Franqueado.")
    p = perfil(db, user)
    if not p["admin"] and not p["franqueados"]:
        raise HTTPException(403, "Seu login ainda não está vinculado a nenhuma unidade. Fale com a franqueadora.")
    return {"token": criar_token(user["usuario"], user["nome"]), **p}


@app.get("/me")
def me(user: dict = UsuarioAtual, db: Session = Depends(get_db)):
    return perfil(db, user)


# ---------------------------------------------------------------- painel

@app.get("/status-opcoes")
def status_opcoes():
    return STATUS_OPCOES


@app.get("/anos")
def anos(user: dict = UsuarioAtual, db: Session = Depends(get_db)):
    q = escopo(db, user, None).with_entities(extract("year", models.Apolice.cancelamento)).distinct()
    lista = sorted({int(a[0]) for a in q}, reverse=True)
    return lista or [date.today().year]


@app.get("/resumo")
def resumo(
    ano: int,
    franqueado: Optional[str] = None,
    user: dict = UsuarioAtual,
    db: Session = Depends(get_db),
):
    q = filtrar(escopo(db, user, franqueado), ano, None, None, None)
    meses = {m: {"mes": m, "qtd": 0, "premio": 0.0, "respondidas": 0} for m in range(1, 13)}
    por_franqueado = defaultdict(lambda: {"qtd": 0, "premio": 0.0, "respondidas": 0, "meses": defaultdict(int)})
    por_status = defaultdict(int)
    for a in q:
        m = meses[a.cancelamento.month]
        m["qtd"] += 1
        m["premio"] += a.premio or 0
        f = por_franqueado[a.franqueado]
        f["qtd"] += 1
        f["premio"] += a.premio or 0
        f["meses"][a.cancelamento.month] += 1
        if a.status:
            m["respondidas"] += 1
            f["respondidas"] += 1
            por_status[a.status] += 1
    total = sum(m["qtd"] for m in meses.values())
    respondidas = sum(m["respondidas"] for m in meses.values())
    return {
        "total": total,
        "premio": round(sum(m["premio"] for m in meses.values()), 2),
        "respondidas": respondidas,
        "pendentes": total - respondidas,
        "meses": [dict(m, premio=round(m["premio"], 2)) for m in meses.values()],
        "por_status": [{"status": s, "qtd": por_status.get(s, 0)} for s in STATUS_OPCOES if por_status.get(s)],
        "franqueados": sorted(
            (
                {
                    "franqueado": nome,
                    "qtd": f["qtd"],
                    "premio": round(f["premio"], 2),
                    "respondidas": f["respondidas"],
                    "meses": {str(k): v for k, v in f["meses"].items()},
                }
                for nome, f in por_franqueado.items()
            ),
            key=lambda x: -x["qtd"],
        ),
    }


@app.get("/apolices")
def listar(
    ano: Optional[int] = None,
    mes: Optional[int] = None,
    franqueado: Optional[str] = None,
    situacao: Optional[str] = None,
    busca: Optional[str] = None,
    user: dict = UsuarioAtual,
    db: Session = Depends(get_db),
):
    q = filtrar(escopo(db, user, franqueado), ano, mes, situacao, busca)
    return [apolice_json(a) for a in q.order_by(models.Apolice.cancelamento, models.Apolice.cliente)]


def _apolice_permitida(db: Session, user: dict, apolice_id: int) -> models.Apolice:
    a = escopo(db, user, None).filter(models.Apolice.id == apolice_id).first()
    if not a:
        raise HTTPException(404, "Apólice não encontrada.")
    return a


@app.get("/apolices/{apolice_id}/respostas")
def respostas(apolice_id: int, user: dict = UsuarioAtual, db: Session = Depends(get_db)):
    _apolice_permitida(db, user, apolice_id)
    rs = (
        db.query(models.Resposta)
        .filter(models.Resposta.apolice_id == apolice_id)
        .order_by(models.Resposta.criado_em.desc())
    )
    return [
        {
            "id": r.id,
            "status": r.status,
            "comentario": r.comentario,
            "autor": r.autor_nome or r.autor_usuario,
            "criado_em": r.criado_em.isoformat() + "Z",
        }
        for r in rs
    ]


class RespostaIn(BaseModel):
    status: str
    comentario: str = ""


@app.post("/apolices/{apolice_id}/respostas")
def responder(apolice_id: int, dados: RespostaIn, user: dict = UsuarioAtual, db: Session = Depends(get_db)):
    if dados.status not in STATUS_OPCOES:
        raise HTTPException(400, "Escolha uma das opções de situação.")
    if dados.status == "Outro motivo" and not dados.comentario.strip():
        raise HTTPException(400, "Descreva o motivo no comentário.")
    a = _apolice_permitida(db, user, apolice_id)
    agora = datetime.utcnow()
    db.add(
        models.Resposta(
            apolice_id=a.id,
            status=dados.status,
            comentario=dados.comentario.strip(),
            autor_usuario=user["usuario"],
            autor_nome=user["nome"],
            criado_em=agora,
        )
    )
    a.status = dados.status
    a.respondido_em = agora
    db.commit()
    return apolice_json(a)


@app.get("/exportar")
def exportar(
    ano: Optional[int] = None,
    mes: Optional[int] = None,
    franqueado: Optional[str] = None,
    situacao: Optional[str] = None,
    busca: Optional[str] = None,
    user: dict = UsuarioAtual,
    db: Session = Depends(get_db),
):
    q = filtrar(escopo(db, user, franqueado), ano, mes, situacao, busca)
    ultimas = {}
    for r in db.query(models.Resposta).order_by(models.Resposta.criado_em):
        ultimas[r.apolice_id] = r
    wb = Workbook()
    ws = wb.active
    ws.title = "Canceladas"
    cab = ["Franqueado", "Cliente", "Seguradora", "Ramo", "Apólice", "Cancelamento", "Início vigência", "Fim vigência",
           "Prêmio líquido", "Motivo", "Comentário", "Respondido por", "Respondido em"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="072A3C")
    for a in q.order_by(models.Apolice.franqueado, models.Apolice.cancelamento):
        r = ultimas.get(a.id)
        ws.append([a.franqueado, a.cliente, a.seguradora, a.produto, a.numero, a.cancelamento, a.vigencia_inicio,
                   a.vigencia_fim, a.premio, a.status or "Pendente", r.comentario if r else "", (r.autor_nome or r.autor_usuario) if r else "",
                   r.criado_em if r else None])
    for col, larg in zip("ABCDEFGHIJKLM", [28, 40, 20, 18, 18, 13, 13, 13, 14, 44, 50, 24, 18]):
        ws.column_dimensions[col].width = larg
    for linha in ws.iter_rows(min_row=2):
        for i in (5, 6, 7):
            linha[i].number_format = "DD/MM/YYYY"
        linha[8].number_format = '"R$" #,##0.00'
        linha[12].number_format = "DD/MM/YYYY HH:MM"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="apolices-canceladas.xlsx"'},
    )


# ---------------------------------------------------------------- admin

def exigir_admin(user: dict = UsuarioAtual, db: Session = Depends(get_db)):
    if not eh_admin(db, user["usuario"]):
        raise HTTPException(403, "Acesso restrito ao gestor.")
    return user


@app.post("/admin/importar")
async def importar(arquivo: UploadFile = File(...), user: dict = Depends(exigir_admin), db: Session = Depends(get_db)):
    conteudo = await arquivo.read()
    try:
        registros, ignoradas = ler_relatorio(conteudo)
    except ValueError as e:
        raise HTTPException(400, str(e))
    except Exception:
        raise HTTPException(400, "Não consegui ler o arquivo. Envie o relatório do Quiver em .XLS ou .XLSX.")
    if not registros:
        raise HTTPException(400, "Nenhuma apólice encontrada na planilha.")

    existentes = {a.chave: a for a in db.query(models.Apolice)}
    novas = atualizadas = 0
    chaves = set()
    agora = datetime.utcnow()
    for reg in registros:
        if reg["chave"] in chaves:
            continue
        chaves.add(reg["chave"])
        a = existentes.get(reg["chave"])
        if a is None:
            db.add(models.Apolice(**reg, importado_em=agora, no_ultimo_relatorio=True))
            novas += 1
        else:
            for campo, valor in reg.items():
                setattr(a, campo, valor)
            a.no_ultimo_relatorio = True
            atualizadas += 1
    # Apólices no mesmo período do relatório que não vieram nele (ex.: cancelamento estornado).
    # Nunca apagamos (as respostas dos franqueados ficam preservadas), só marcamos.
    inicio = min(r["cancelamento"] for r in registros)
    fim = max(r["cancelamento"] for r in registros)
    for a in existentes.values():
        if a.chave not in chaves and inicio <= a.cancelamento <= fim:
            a.no_ultimo_relatorio = False
    db.add(models.Importacao(arquivo=arquivo.filename or "", autor=user["nome"], linhas=len(registros),
                             novas=novas, atualizadas=atualizadas, ignoradas=ignoradas, criado_em=agora))
    db.commit()
    return {"linhas": len(registros), "novas": novas, "atualizadas": atualizadas, "ignoradas": ignoradas,
            "periodo": [inicio.isoformat(), fim.isoformat()]}


@app.get("/admin/importacoes")
def importacoes(user: dict = Depends(exigir_admin), db: Session = Depends(get_db)):
    return [
        {"id": i.id, "arquivo": i.arquivo, "autor": i.autor, "linhas": i.linhas, "novas": i.novas,
         "atualizadas": i.atualizadas, "ignoradas": i.ignoradas, "criado_em": i.criado_em.isoformat() + "Z"}
        for i in db.query(models.Importacao).order_by(models.Importacao.criado_em.desc()).limit(30)
    ]


@app.get("/admin/acessos")
def acessos(user: dict = Depends(exigir_admin), db: Session = Depends(get_db)):
    contagem = dict(
        db.query(models.Apolice.franqueado, func.count(models.Apolice.id)).group_by(models.Apolice.franqueado).all()
    )
    vinc = defaultdict(list)
    for v in db.query(models.Vinculo):
        vinc[v.franqueado].append(v.usuario)
    nomes = sorted(set(contagem) | set(vinc))
    try:
        portal = [{"usuario": str(u.get("u", "")).lower(), "nome": u.get("nome", "")} for u in usuarios_portal()]
    except HTTPException:
        portal = []
    vistos, usuarios = set(), []
    for u in sorted(portal, key=lambda x: x["nome"].lower()):
        if u["usuario"] and u["usuario"] not in vistos:
            vistos.add(u["usuario"])
            usuarios.append(u)
    return {
        "franqueados": [{"franqueado": n, "apolices": contagem.get(n, 0), "usuarios": sorted(vinc.get(n, []))} for n in nomes],
        "usuarios_portal": usuarios,
        "admins": sorted(a.usuario for a in db.query(models.Admin)),
    }


class VinculoIn(BaseModel):
    franqueado: str
    usuarios: List[str]


@app.put("/admin/vinculos")
def salvar_vinculos(dados: VinculoIn, user: dict = Depends(exigir_admin), db: Session = Depends(get_db)):
    db.query(models.Vinculo).filter(models.Vinculo.franqueado == dados.franqueado).delete()
    for u in sorted({u.strip().lower() for u in dados.usuarios if u.strip()}):
        db.add(models.Vinculo(usuario=u, franqueado=dados.franqueado))
    db.commit()
    return {"ok": True}


class AdminsIn(BaseModel):
    usuarios: List[str]


@app.put("/admin/admins")
def salvar_admins(dados: AdminsIn, user: dict = Depends(exigir_admin), db: Session = Depends(get_db)):
    novos = {u.strip().lower() for u in dados.usuarios if u.strip()}
    if user["usuario"] not in novos:
        raise HTTPException(400, "Você não pode remover o seu próprio acesso de gestor.")
    db.query(models.Admin).delete()
    for u in sorted(novos):
        db.add(models.Admin(usuario=u))
    db.commit()
    return {"ok": True}


@app.get("/admin/backup")
def backup(user: dict = Depends(exigir_admin), db: Session = Depends(get_db)):
    """Cópia de segurança do que só existe aqui (respostas, vínculos, gestores) — a planilha pode ser reimportada."""
    return {
        "gerado_em": datetime.utcnow().isoformat() + "Z",
        "respostas": [
            {"chave": a.chave, "status": r.status, "comentario": r.comentario, "autor_usuario": r.autor_usuario,
             "autor_nome": r.autor_nome, "criado_em": r.criado_em.isoformat()}
            for r, a in db.query(models.Resposta, models.Apolice).join(models.Apolice, models.Apolice.id == models.Resposta.apolice_id)
        ],
        "vinculos": [{"usuario": v.usuario, "franqueado": v.franqueado} for v in db.query(models.Vinculo)],
        "admins": [a.usuario for a in db.query(models.Admin)],
    }
